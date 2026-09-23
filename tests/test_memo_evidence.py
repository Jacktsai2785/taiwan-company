import unittest
from unittest.mock import AsyncMock, patch

from services import memo_extractor


class MemoEvidenceTests(unittest.IsolatedAsyncioTestCase):
    def test_material_candidate_scan_keeps_numeric_and_governance_context(self):
        transcript = "\n".join([
            "[00:00:01] 普通開場",
            "[00:00:02] 訂單從8000元提高到8萬元",
            "[00:00:03] 這是成長轉折",
            "[00:00:04] 普通過場",
            "[00:00:05] 未來要設計股權與表決權",
        ])

        candidates = memo_extractor._material_candidate_text(transcript)

        self.assertIn("訂單從8000元提高到8萬元", candidates)
        self.assertIn("未來要設計股權與表決權", candidates)

    def test_deterministic_safety_net_preserves_high_signal_verbatim(self):
        transcript = "\n".join([
            "[00:00:01] 客戶訂單從8000元",
            "[00:00:02] 提高到至少8萬元",
            "[00:00:03] 未來要設計股權與表決權",
        ])

        evidence = memo_extractor._deterministic_material_evidence(transcript)

        self.assertTrue(any("8000元" in item["quote"] for item in evidence["financials"]))
        self.assertTrue(any("股權" in item["quote"] for item in evidence["investment_terms"]))
        self.assertEqual(evidence["financials"][0]["fact"], "客戶訂單從8000元")

    def test_generic_signals_cover_saas_manufacturing_and_biotech(self):
        transcript = "\n".join([
            "[00:00:01] 我們ARR為1,200萬元，年增35%",
            "[00:00:02] 目前團隊42人",
            "[00:00:03] 廠房月產10萬件，良率99.2%",
            "[00:00:04] 預計2028年完成三期臨床",
            "[00:00:05] 本輪募資US$5 million，釋股12%",
        ])

        evidence = memo_extractor._deterministic_material_evidence(transcript)

        self.assertTrue(evidence["financials"])
        self.assertTrue(evidence["headcount"])
        self.assertTrue(evidence["factory_capacity"])
        self.assertTrue(evidence["recent_development"])
        self.assertTrue(evidence["investment_terms"])

    def test_generic_safety_net_excludes_questions_and_unrelated_examples(self):
        transcript = "\n".join([
            "[00:00:01] 你們營收有沒有1億元",
            "[00:00:02] 今年營收是3億元嗎？",
            "[00:00:03] 假設做到Uber要幾十億美元",
            "[00:00:04] 這是別人的案例",
        ])

        evidence = memo_extractor._deterministic_material_evidence(transcript)

        self.assertEqual(evidence["financials"], [])

    async def test_dual_extraction_keeps_success_when_other_path_times_out(self):
        successful = {key: [] for key in memo_extractor._EXTRACT_KEYS}
        successful["financials"] = [
            {"fact": "營收成長", "quote": "營收成長", "timestamp": ""}
        ]
        with patch(
            "services.memo_extractor._extract_evidence_once",
            new=AsyncMock(side_effect=RuntimeError("timeout")),
        ), patch(
            "services.memo_extractor._extract_material_facts_once",
            new=AsyncMock(return_value=successful),
        ):
            parts = await memo_extractor._extract_chunk_dual(
                "測試公司", "營收成長", "codex"
            )

        self.assertEqual(parts, [successful])

    async def test_candidate_audit_failure_does_not_discard_dual_extraction_results(self):
        # The candidate-chunk omission audit is supplementary; one failing
        # chunk must not wipe out results the main dual-path extraction
        # already succeeded at getting.
        transcript = "營收成長至十億元。"
        successful = {key: [] for key in memo_extractor._EXTRACT_KEYS}
        successful["financials"] = [
            {"fact": "營收成長", "quote": "營收成長", "timestamp": ""}
        ]
        with patch(
            "services.memo_extractor._extract_evidence_once",
            new=AsyncMock(return_value=successful),
        ), patch(
            "services.memo_extractor._extract_material_facts_once",
            new=AsyncMock(side_effect=RuntimeError("candidate audit timeout")),
        ), patch(
            "services.memo_extractor._material_candidate_text",
            return_value=transcript,
        ), patch(
            "services.memo_extractor._synthesize_fields_from_evidence",
            new=AsyncMock(return_value=(
                {key: "" for key in memo_extractor._EXTRACT_KEYS},
                {key: {"evidence_count": 0, "used_count": 0, "synthesized": False}
                 for key in memo_extractor._EXTRACT_KEYS},
            )),
        ):
            result, audit = await memo_extractor.extract_with_audit(
                "測試公司", transcript, engine="codex"
            )

        self.assertIn(
            "營收成長",
            [item["fact"] for item in audit["evidence"]["financials"]],
        )

    def test_rejects_fact_when_quote_is_not_in_transcript(self):
        claimed = {
            "financials": [{
                "fact": "營收十億元",
                "quote": "營收十億元",
                "timestamp": "",
            }]
        }

        evidence = memo_extractor._validated_evidence(
            claimed, "逐字稿只說產品已經上市。"
        )

        self.assertEqual(evidence["financials"], [])

    async def test_all_chunks_are_aggregated_without_synthesis_loss(self):
        transcript = "甲" * memo_extractor._CHUNK_CHARS + "乙" * 600
        first = {key: [] for key in memo_extractor._EXTRACT_KEYS}
        second = {key: [] for key in memo_extractor._EXTRACT_KEYS}
        first["business_revenue"] = [
            {"fact": "提供甲產品", "quote": "甲", "timestamp": ""}
        ]
        second["business_revenue"] = [
            {"fact": "提供乙服務", "quote": "乙", "timestamp": ""}
        ]
        synthesized = {key: "" for key in memo_extractor._EXTRACT_KEYS}
        synthesized["business_revenue"] = "提供甲產品，並提供乙服務。"
        coverage = {
            key: {"evidence_count": 0, "used_count": 0, "synthesized": False}
            for key in memo_extractor._EXTRACT_KEYS
        }
        coverage["business_revenue"] = {
            "evidence_count": 2, "used_count": 2, "synthesized": True,
        }
        with patch(
            "services.memo_extractor._extract_evidence_once",
            new=AsyncMock(side_effect=[first, second]),
        ), patch(
            "services.memo_extractor._extract_material_facts_once",
            new=AsyncMock(side_effect=[
                {key: [] for key in memo_extractor._EXTRACT_KEYS},
                {key: [] for key in memo_extractor._EXTRACT_KEYS},
            ]),
        ), patch(
            "services.memo_extractor._synthesize_fields_from_evidence",
            new=AsyncMock(return_value=(synthesized, coverage)),
        ):
            result, audit = await memo_extractor.extract_with_audit(
                "測試公司", transcript, engine="codex"
            )

        evidence = audit["evidence"]
        self.assertEqual(len(evidence["business_revenue"]), 2)
        self.assertEqual(result["business_revenue"], "提供甲產品，並提供乙服務。")
        self.assertTrue(audit["coverage"]["business_revenue"]["synthesized"])
        self.assertEqual(
            [item["id"] for item in evidence["business_revenue"]],
            ["business_revenue:1", "business_revenue:2"],
        )

    def test_safety_net_is_audit_only_when_model_evidence_exists(self):
        items = [
            {"id": "headcount:1", "fact": "公司目前約80人。", "quote": "80人"},
            {
                "id": "headcount:2",
                "fact": "八十个人开口等你养",
                "quote": "八十个人开口等你养",
                "source": "deterministic_safety_net",
            },
        ]

        selected = memo_extractor._deduplicated_facts(items)

        self.assertEqual([item["id"] for item in selected], ["headcount:1"])

    def test_safety_net_never_becomes_published_prose_by_itself(self):
        items = [{
            "id": "factory_capacity:1",
            "fact": "其他半導體公司的產能有限",
            "quote": "其他半導體公司的產能有限",
            "source": "deterministic_safety_net",
        }]

        self.assertEqual(memo_extractor._deduplicated_facts(items), [])
        self.assertEqual(memo_extractor._fallback_field_text(items), "")

    def test_prose_normalization_removes_broken_punctuation(self):
        self.assertEqual(
            memo_extractor._normalize_memo_prose("第一句。；第二句；。第三句。。"),
            "第一句。第二句。第三句。",
        )

    async def test_synthesis_requires_valid_evidence_ids(self):
        evidence = {key: [] for key in memo_extractor._EXTRACT_KEYS}
        evidence["business_revenue"] = [
            {
                "id": "business_revenue:1",
                "fact": "公司提供企業軟體開發服務。",
                "quote": "提供企業軟體開發服務",
                "timestamp": "00:01:00",
            }
        ]
        response = {
            "business_revenue": {
                "text": "公司提供企業軟體開發服務。",
                "evidence_ids": ["business_revenue:1"],
            }
        }
        with patch(
            "services.memo_extractor.claude_client.ask",
            return_value=__import__("json").dumps(response, ensure_ascii=False),
        ), patch(
            "services.memo_extractor.asyncio.to_thread",
            new=AsyncMock(return_value=__import__("json").dumps(response, ensure_ascii=False)),
        ):
            result = await memo_extractor._synthesize_field_group(
                "測試公司", ("business_revenue",), evidence, "codex"
            )

        self.assertEqual(result, response)

    async def test_synthesize_financials_table_fills_only_supported_fields(self):
        """財務狀況儲存格裡的巢狀表格（年度 x Now/Now+1/Now+2/Now+3）要能被 AI
        自動填數字進去；evidence 沒明確支持的期別/指標要留空，不能亂填或推算。
        數字欄位不該留「元」這種單位字樣（表格已經是數字欄，不需要單位），
        毛利率(%) 是唯一例外，要保留 %。"""
        evidence = [{
            "id": "financials:1",
            "fact": "依114年08月31日暫結資產負債表，民114年1月至8月營業收入淨額為14,452,108元，毛利率71.02%",
        }]
        response = {"now": {
            "period_label": "114/8/31",
            "revenue": "14,452,108元",
            "gross_margin_pct": "71.02%",
        }}
        with patch(
            "services.memo_extractor.asyncio.to_thread",
            new=AsyncMock(return_value=__import__("json").dumps(response, ensure_ascii=False)),
        ):
            result = await memo_extractor._synthesize_financials_table(
                "測試公司", evidence, "claude"
            )

        self.assertEqual(result["fin_now_revenue"], "14,452,108")  # 「元」被去掉
        self.assertEqual(result["fin_now_gross_margin_pct"], "71.02%")  # % 保留
        self.assertEqual(result["fin_now_cogs"], "")
        self.assertEqual(result["fin_now1_revenue"], "")
        self.assertEqual(result["fin_period_label_now"], "114/8/31")
        self.assertEqual(result["fin_period_label_now1"], "")
        self.assertEqual(
            len(result),
            len(memo_extractor.FINANCIAL_TABLE_KEYS) + len(memo_extractor.FINANCIAL_PERIOD_LABEL_KEYS),
        )

    async def test_synthesize_financials_table_skips_ai_call_without_evidence(self):
        with patch(
            "services.memo_extractor.asyncio.to_thread",
            new=AsyncMock(side_effect=AssertionError("should not call AI without evidence")),
        ):
            result = await memo_extractor._synthesize_financials_table("測試公司", [], "claude")

        self.assertTrue(all(v == "" for v in result.values()))

    def test_strip_currency_unit_removes_yuan_but_keeps_digits(self):
        self.assertEqual(memo_extractor._strip_currency_unit("14,452,108元"), "14,452,108")
        self.assertEqual(memo_extractor._strip_currency_unit("5,000萬元"), "5,000萬")
        self.assertEqual(memo_extractor._strip_currency_unit("71.02%"), "71.02%")

    async def test_extract_evidence_dual_from_file_survives_one_pass_failing(self):
        """比照文字管線的雙抽取設計：讀圖有兩條獨立 pass，任一條掛掉或漏抓
        不該讓另一條的結果也不見。"""
        ok_result = {key: [] for key in memo_extractor._EXTRACT_KEYS}
        ok_result["management_team"] = [{"fact": "林妙娟為創辦人", "quote": "", "timestamp": "", "source": "vision_extraction"}]
        with patch(
            "services.memo_extractor._extract_evidence_from_file",
            new=AsyncMock(return_value=ok_result),
        ), patch(
            "services.memo_extractor._extract_material_facts_from_file",
            new=AsyncMock(side_effect=RuntimeError("模型逾時")),
        ):
            results = await memo_extractor._extract_evidence_dual_from_file(
                "測試公司", "/tmp/fake.pdf", "claude"
            )

        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["management_team"][0]["fact"], "林妙娟為創辦人")

    def test_fill_template_fills_nested_financials_table(self):
        """財務狀況儲存格裡本來就有一張巢狀表格（年度 x Now~Now+3），fill_template
        要能把 fin_<period>_<metric> 的值填進對應格子，不能只塞長文字進儲存格。"""
        from docx import Document
        import io

        memo = {
            "fin_now_revenue": "14,452,108",
            "fin_now_gross_margin_pct": "71.02%",
            "fin_now_operating_income": "5,000,000",
            "fin_now_non_operating_income": "100,000",
            "fin_now_non_operating_expense": "50,000",
            "fin_period_label_now": "114/8/31",
        }
        docx_bytes = memo_extractor.fill_template({"name": "測試公司"}, memo, "2026/09/23")

        doc = Document(io.BytesIO(docx_bytes))
        fin_cell = doc.tables[0].rows[10].cells[0]
        nested = fin_cell.tables[0]
        rows_by_label = {row.cells[0].text.strip(): [c.text.strip() for c in row.cells] for row in nested.rows}

        self.assertEqual(rows_by_label["營收"][1], "14,452,108")
        self.assertEqual(rows_by_label["毛利率(%)"][1], "71.02%")
        self.assertEqual(rows_by_label["COGS"][1], "")  # 沒 evidence 支持，留空
        # 費用列已更名為「營業費用」，且範本裡插入了 G&A人數 + 營業利益/業外收入/業外支出
        self.assertIn("營業費用", rows_by_label)
        self.assertNotIn("費用", rows_by_label)
        self.assertIn("G&A人數", rows_by_label)
        self.assertEqual(rows_by_label["營業利益"][1], "5,000,000")
        self.assertEqual(rows_by_label["業外收入"][1], "100,000")
        self.assertEqual(rows_by_label["業外支出"][1], "50,000")
        # 有實際財報期間時，表頭「Now」要換成真正的期間文字
        self.assertEqual(rows_by_label["年度"][1], "114/8/31")
        # 順序：G&A人數 在 G&A 跟 R&D 之間；RD人數 之後、稅後淨利 之前是三個新指標
        labels = list(rows_by_label.keys())
        self.assertLess(labels.index("G&A"), labels.index("G&A人數"))
        self.assertLess(labels.index("G&A人數"), labels.index("R&D"))
        self.assertLess(labels.index("RD人數"), labels.index("營業利益"))
        self.assertLess(labels.index("業外支出"), labels.index("稅後淨利"))


if __name__ == "__main__":
    unittest.main()
